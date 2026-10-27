"""Værkets tidsserier i motoren (KONTRAKT v2.3 §6d, brief 18). Hver test fejler før runde 5.

Fire lag: den normaliserede fil (``fjv.tidsserie``), frysningen -> case + argv
(``byg_case``), arbejderens hentning med hash-kontrol (``koersel.hent_tidsserie`` og
hele ``koer_en`` med falsk runtime), og de betroede skrivninger (repo-registrering,
opløsning af aktiv/valgt, værksark-import) mod forbindelsesstubben fra test_k2.

Modellens egen læser (``apply_heat_csv_override``/profil-læseren) bruges, hvor
pandas findes (``_lokal/.venv``), til at bevise at den normaliserede fil giver
**bitvis** samme varmelast og solprofil som repoets kildefil.
"""

from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

HER = Path(__file__).resolve().parent
sys.path.insert(0, str(HER.parent))
from fjv import datakrav as DK  # noqa: E402
from fjv import db as DB  # noqa: E402
from fjv import koersel as KS  # noqa: E402
from fjv import konstanter as K  # noqa: E402
from fjv import runtime as R  # noqa: E402
from fjv import tidsserie as TS  # noqa: E402
from fjv import tidsserie_db as TDB  # noqa: E402
from fjv import vaerksark as VS  # noqa: E402
from fjv.byg_case import InputFejl, byg, input_hash, raekker_fra_vaerk  # noqa: E402
from fjv.flad import flad  # noqa: E402
from tests.test_k2 import StubDb  # noqa: E402


def kald(db, ord_):
    """(sql, params) for de logførte sætninger, der indeholder ``ord_``."""
    return [(s, p) for s, p in db.log if ord_ in s]

MODELREPO = Path(__file__).resolve().parents[3] / "fjernvarme-businesscase"
KODE = HER.parent / "_lokal" / "Files" / K.REL_KODE / K.MODEL_COMMIT
REPO_DATA = next((p for p in (KODE / "data", MODELREPO / "data") if p.is_dir()), None)
UUID1 = "11111111-2222-4333-8444-555555555555"
UUID2 = "aaaaaaaa-bbbb-4ccc-8ddd-eeeeeeeeeeee"

try:
    import pandas  # noqa: F401
    HAR_PANDAS = True
except ImportError:
    HAR_PANDAS = False

CASE = """\
time:
  start: "2026-01-05T00:00:00Z"
  end: "2026-01-06T23:00:00Z"
  resolution: "1h"
prices:
  straw:
    value: 232.0
units:
  halm:
    type: biomass_boiler
    p_max_heat: 12.0
  sol:
    type: solar_thermal
    p_max_heat: 8.0
    production_profile_path: "data/solar_andeby_dmi_karup.csv"
"""


def timer(fra="2026-01-05", doegn=3, mw=lambda t: 10 + t % 5, spring=()):
    """En gyldig timeserie (oraklet afviser under 48 timer)."""
    from datetime import date, timedelta
    d0 = date.fromisoformat(fra)
    ud = []
    for i in range(doegn * 24):
        if i in spring:
            continue
        d = d0 + timedelta(days=i // 24)
        ud.append(f"{d.isoformat()}T{i % 24:02d}:00:00Z,{mw(i)}\n")
    return "time,mw\n" + "".join(ud)


def serie(linjer: str, rolle="varmelast", **kw):
    """Normaliseret serie via fv-data's orakel - præcis den vej motoren bruger."""
    return DK.valider(linjer, rolle, **kw)


def raekker(case=CASE, aendringer=(), **scen):
    scen.setdefault("model_commit", K.MODEL_COMMIT)
    scen.setdefault("data_commit", K.DF_DATA_SHA)
    return raekker_fra_vaerk(flad(case), aendringer, scen)


def ref(sha="a" * 64, id_=UUID1):
    return {"id": id_, "sha256": sha}


# ------------------------------------------------------------------ den normaliserede fil

class TestNormalisering(unittest.TestCase):
    """Motorens brug af oraklet + dens egen strenge læsning af filen. Reglerne selv
    (fundkoder, tærskler) testes af fv-data (viden/vaerktoej/tjek_data_krav.py)."""

    def test_form_og_hash(self):
        rk = "".join(f"2025-03-{d:02d} {h:02d}:00:00,{10 + h % 5}\n" for d in (1, 2, 3) for h in range(24))
        r = serie("timestamp,heat_mw_abvaerk\n2025-02-28 12:00:00,19.74\n2025-02-28 13:00:00,18.2\n" + rk)
        self.assertEqual(r.status, "gyldig", r.validering.get("fejl"))
        self.assertTrue(r.data.startswith(b"time,mw\n2025-02-28T12:00:00Z,19.74\n2025-02-28T13:00:00Z,18.2\n"))
        self.assertEqual(r.sha256, TS.sha256(r.data))
        pkt = TS.laes_normaliseret(r.data)
        self.assertEqual(len(pkt), r.antal_raekker)
        self.assertEqual((r.start, TS.tid_tekst(pkt[0][0])), ("2025-02-28T12:00:00Z", "2025-02-28T12:00:00Z"))

    def test_kildens_taltekst_bevares_tegn_for_tegn(self):
        # pandas' hurtige parser lover ikke samme double for "1.234560" og "1.23456".
        rk = "".join(f"2025-06-{d:02d}T{h:02d}:00:00Z,{'1.234560' if 9 <= h <= 13 else '0.000000'}\n"
                     for d in (1, 2, 3) for h in range(24))
        r = serie("time,power_mw\n" + rk, "solprofil")
        self.assertEqual(r.status, "gyldig", r.validering.get("fejl"))
        self.assertIn(b",0.000000\n", r.data)
        self.assertIn(b",1.234560\n", r.data)

    def test_orakel_og_kilder_er_registreret(self):
        k = DK.kilder()
        self.assertNotEqual(k["valider_tidsserie.py"], "MANGLER")
        self.assertEqual(k["version"], DK.krav()["version"])

    def test_tal_tekst_vektorer_fra_data_krav(self):
        for x, forventet in DK.krav()["tidsserier"]["normaliseret_fil"]["tal_tekst_testvektorer"]:
            self.assertEqual(DK.orakel().tal_tekst(float(x)), forventet)

    def test_laes_normaliseret_afviser_andre_former(self):
        for data in (b"timestamp,mw\n2025-01-01T00:00:00Z,1\n",      # forkert overskrift
                     b"time,mw\n2025-01-01 00:00:00,1\n",            # ikke ISO Z
                     b"time,mw\n2025-01-01T00:30:00Z,1\n",           # ikke hel time
                     b"time,mw\n2025-01-01T00:00:00Z,1",             # afkortet
                     b"time,mw\n2025-01-01T01:00:00Z,1\n2025-01-01T00:00:00Z,1\n",
                     b"time,mw\n2025-01-01T00:00:00Z,1\n2025-01-01T00:00:00Z,2\n",
                     b"time,mw\n2025-01-01T00:00:00Z,\n",
                     b"time,mw\n2025-01-01T00:00:00Z,-1\n",
                     b"time,mw\n2025-01-01T00:00:00Z,1e-05\n",
                     b"time,mw\n"):
            with self.assertRaises(TS.TidsserieFejl, msg=data):
                TS.laes_normaliseret(data)

    def test_lokalt_navn_og_lakehouse_sti(self):
        self.assertEqual(TS.lokalt_navn("ab" * 32), "_fjv_ts_abababababababab.csv")
        self.assertEqual(TS.lakehouse_rel("andeby", "c" * 64), f"fjv/tidsserier/andeby/{'c' * 64}.csv")
        with self.assertRaises(TS.TidsserieFejl):
            TS.lokalt_navn("../../etc")


@unittest.skipUnless(HAR_PANDAS and MODELREPO.is_dir(), "kræver modelrepoet og pandas (_lokal/.venv)")
class TestModellensLaeserErBitvisEns(unittest.TestCase):
    """Den normaliserede fil giver PRÆCIS samme serie i modellen som repoets fil."""

    @classmethod
    def setUpClass(cls):
        sys.path.insert(0, str(MODELREPO))
        cls.tmp = Path(tempfile.mkdtemp())

    def _norm(self, navn, rolle):
        r = DK.valider((MODELREPO / "data" / navn).read_text(encoding="utf-8"), rolle)
        p = self.tmp / f"{r.sha256}.csv"
        p.write_bytes(r.data)
        return p

    def test_varmelast_splejset(self):
        import numpy as np
        import pandas as pd
        import xarray as xr
        from src.data_loader import apply_heat_csv_override
        norm = self._norm("billund_abvaerk_hourly_splejset_jun2026.csv", "varmelast")
        ds = xr.Dataset(coords={"time": pd.date_range("2025-03-01", "2026-06-30 23:00", freq="h")})
        a = apply_heat_csv_override(ds, MODELREPO / "data" / "billund_abvaerk_hourly_splejset_jun2026.csv",
                                    column="heat_mw_abvaerk")["heat_demand"].values
        b = apply_heat_csv_override(ds, norm, column="mw")["heat_demand"].values
        self.assertTrue(np.array_equal(a.view(np.uint64), b.view(np.uint64)))

    def test_solprofil_andeby(self):
        import numpy as np
        import pandas as pd

        def laes(p):        # samme læsning som data_loader._attach_unit_profiles
            df = pd.read_csv(p, parse_dates=["time"])
            df["time"] = pd.to_datetime(df["time"], utc=True).dt.tz_convert(None)
            return df.set_index("time").sort_index().iloc[:, 0]
        a = laes(MODELREPO / "data" / "solar_andeby_dmi_karup.csv")
        b = laes(self._norm("solar_andeby_dmi_karup.csv", "solprofil"))
        self.assertTrue(a.index.equals(b.index))
        self.assertTrue(np.array_equal(a.values.view(np.uint64), b.values.view(np.uint64)))


# ------------------------------------------------------------------ frysning -> case + argv

class TestBygMedTidsserier(unittest.TestCase):
    def test_varmelast_bliver_heat_csv_med_kolonnen_mw(self):
        sha = "b" * 64
        plan = byg(raekker(tidsserier={"varmelast": ref(sha)}))
        self.assertEqual(plan.tidsserier, {"varmelast": ref(sha)})
        i = plan.cli.index("--heat-csv")
        self.assertEqual(plan.cli[i:i + 4], ["--heat-csv", f"data/{TS.lokalt_navn(sha)}",
                                             "--heat-csv-column", "mw"])

    def test_varmelast_kraever_filen_i_data_ved_anden_byg(self):
        tmp = Path(tempfile.mkdtemp())
        (tmp / "solar_andeby_dmi_karup.csv").write_text("x", encoding="utf-8")
        r = raekker(tidsserier={"varmelast": ref("b" * 64)})
        with self.assertRaises(InputFejl):
            byg(r, heat_csv_rod=tmp)
        (tmp / TS.lokalt_navn("b" * 64)).write_text("x", encoding="utf-8")
        self.assertIn("--heat-csv", byg(r, heat_csv_rod=tmp).cli)
        shutil.rmtree(tmp)

    def test_syntetisk_og_udeladt_giver_ingen_heat_csv(self):
        for scen in ({"tidsserier": {"varmelast": "syntetisk"}}, {}):
            plan = byg(raekker(**scen))
            self.assertNotIn("--heat-csv", plan.cli)
            self.assertEqual(plan.tidsserier, {})

    def test_heat_csv_og_tidsserie_er_dublet(self):
        with self.assertRaises(InputFejl):
            byg(raekker(heat_csv="billund_abvaerk_hourly.csv", tidsserier={"varmelast": ref()}))

    def test_solprofil_erstatter_den_ene_solenhed(self):
        sha = "c" * 64
        plan = byg(raekker(tidsserier={"solprofil": ref(sha)}))
        self.assertEqual(plan.case["units"]["sol"]["production_profile_path"],
                         f"data/{TS.lokalt_navn(sha)}")
        self.assertIn(f"data/{TS.lokalt_navn(sha)}", plan.case_yaml)

    def test_markoer_kraever_frossen_solprofil(self):
        case = CASE.replace('"data/solar_andeby_dmi_karup.csv"', f'"{TS.MARKOER_SOLPROFIL}"')
        with self.assertRaises(InputFejl) as f:
            byg(raekker(case))
        self.assertIn("solprofil-tidsserie", str(f.exception))
        plan = byg(raekker(case, tidsserier={"solprofil": ref("d" * 64)}))
        self.assertEqual(plan.case["units"]["sol"]["production_profile_path"],
                         f"data/{TS.lokalt_navn('d' * 64)}")

    def test_markoer_paa_slaaet_fra_enhed_er_ok_uden_solprofil(self):
        case = CASE.replace('"data/solar_andeby_dmi_karup.csv"', f'"{TS.MARKOER_SOLPROFIL}"')
        plan = byg(raekker(case, [{"type": "disable", "sti": "sol"}]))
        self.assertIsNone(plan.case["units"]["sol"]["production_profile_path"])

    def test_solprofil_og_aendret_profilsti_er_dublet(self):
        a = [{"type": "set", "sti": "units.sol.production_profile_path",
              "vaerdi_json": '"data/synthetic_solar_andeby.csv"'}]
        with self.assertRaises(InputFejl):
            byg(raekker(aendringer=a, tidsserier={"solprofil": ref()}))

    def test_to_solenheder_uden_markoer_er_tvetydigt(self):
        case = CASE + "  sol2:\n    type: solar_thermal\n    p_max_heat: 2.0\n"
        with self.assertRaises(InputFejl):
            byg(raekker(case, tidsserier={"solprofil": ref()}))

    def test_ingen_aktiv_solenhed_er_advarsel(self):
        plan = byg(raekker(aendringer=[{"type": "disable", "sti": "sol"}],
                           tidsserier={"solprofil": ref()}))
        self.assertTrue(any("ingen aktiv solvarmeenhed" in a for a in plan.advarsler))

    def test_ugyldige_referencer_afvises(self):
        P = "_koersel.tidsserie."
        basis = raekker()
        for sti, v in ((P + "varmelast", {"id": "x", "sha256": "a" * 64}),
                       (P + "varmelast", {"id": UUID1, "sha256": "A" * 64}),
                       (P + "varmelast", {"id": UUID1}),
                       (P + "solprofil", "syntetisk"),
                       (P + "vind", ref())):
            r = basis + [{"kilde": "scenarie", "sti": sti, "vaerdi_json": json.dumps(v)}]
            with self.assertRaises(InputFejl, msg=(sti, v)):
                byg(r)

    def test_frysningen_er_i_hashen(self):
        a = raekker(tidsserier={"varmelast": ref("e" * 64)})
        b = raekker(tidsserier={"varmelast": ref("f" * 64)})
        self.assertNotEqual(input_hash(a), input_hash(b))
        self.assertIn({"kilde": "scenarie", "sti": "_koersel.tidsserie.varmelast",
                       "vaerdi_json": json.dumps(ref("e" * 64), separators=(",", ":"))}, a)


# ------------------------------------------------------------------ arbejderen henter

class TestHentTidsserie(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.gl = K.FILES
        K.FILES = str(self.tmp)
        self.r = serie(timer())
        TDB.gem_tidsseriefil("andeby", self.r.data, self.r.sha256)

    def tearDown(self):
        K.FILES = self.gl
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_god_fil_laegges_i_data(self):
        ud = KS.hent_tidsserie(None, "varmelast", ref(self.r.sha256), self.tmp / "kode" / "data")
        self.assertIsInstance(ud, dict, ud)
        self.assertEqual((self.tmp / "kode" / "data" / TS.lokalt_navn(self.r.sha256)).read_bytes(),
                         self.r.data)

    def test_registreret_sti_bruges(self):
        class L:
            def tidsserie_sti(self, i):
                return f"Files/fjv/tidsserier/andeby/{sha}.csv"
        sha = self.r.sha256
        ud = KS.hent_tidsserie(L(), "varmelast", ref(sha), self.tmp / "d")
        self.assertTrue(ud["fil"].endswith(f"andeby/{sha}.csv"))

    def test_aendret_fil_afvises(self):
        p = self.tmp / TS.lakehouse_rel("andeby", self.r.sha256)
        p.write_bytes(self.r.data.replace(b"T01:00:00Z,11\n", b"T01:00:00Z,12\n", 1))
        ud = KS.hent_tidsserie(None, "varmelast", ref(self.r.sha256), self.tmp / "d")
        self.assertIsInstance(ud, str)
        self.assertIn("ændret efter bestillingen", ud)
        self.assertFalse((self.tmp / "d").exists())

    def test_manglende_fil(self):
        ud = KS.hent_tidsserie(None, "solprofil", ref("0" * 64), self.tmp / "d")
        self.assertIn("findes ikke", ud)

    def test_uforanderlig_fil_med_andet_indhold_er_fejl(self):
        p = self.tmp / TS.lakehouse_rel("andeby", self.r.sha256)
        p.write_bytes(b"noget andet")
        with self.assertRaises(RuntimeError):
            TDB.gem_tidsseriefil("andeby", self.r.data, self.r.sha256)


class TestKoerEnMedTidsserie(unittest.TestCase):
    """Hele koer_en med falsk runtime: den frosne varmelast ender som --heat-csv på en fil,
    der FINDES i kørslens data/ med præcis de bestilte bytes."""

    def test_argv_og_fil(self):
        from tests.test_k3 import HER as K3
        tmp = Path(tempfile.mkdtemp())
        gl = K.FILES
        K.FILES = str(tmp)
        kode = tmp / "kode"
        (kode / "data").mkdir(parents=True)
        (kode / "data" / "solar_andeby_dmi_karup.csv").write_text("x", encoding="utf-8")
        (kode / "requirements.txt").write_text("x\n", encoding="utf-8")
        fixture = K3 / "data" / "f1_andeby"
        r = serie(timer())
        TDB.gem_tidsseriefil("andeby", r.data, r.sha256, files=str(tmp))
        set_ = {}

        def falsk_proces(argv, cwd, log, timeout_s, ved_linje=None, stop=None):
            set_["argv"] = list(argv)
            i = argv.index("--heat-csv")
            set_["fil"] = (Path(cwd) / argv[i + 1]).read_bytes()
            ud = Path(argv[argv.index("--out-dir") + 1])
            for f in fixture.iterdir():
                if f.name != "log.txt":
                    shutil.copy(f, ud / f.name)
            shutil.copy(fixture / "log.txt", log)
            return 0

        erstat = {"sikr_model": lambda sha, files=None: (kode, {}),
                  "sikr_dfdata": lambda sha, files=None: (tmp / "df", {}),
                  "lokal_kodekopi": lambda kilde, maal, dfdata=None, git_init=False: kode,
                  "byg_miljoe": lambda req, venv, log: "python",
                  "pakkeversioner": lambda py, navne: {},
                  "maal_runtime": lambda: {"python": "3.12"},
                  "koer_proces": falsk_proces}
        gamle = {n: getattr(R, n) for n in erstat}
        try:
            for n, f in erstat.items():
                setattr(R, n, f)
            from fjv.rapport import Rapport
            rk = raekker(tidsserier={"varmelast": ref(r.sha256)})
            res = KS.koer_en(rk, "ts1", KS.IngenLager(), Rapport("t", sti=str(tmp / "r.json")))
        finally:
            for n, f in gamle.items():
                setattr(R, n, f)
            K.FILES = gl
        self.assertEqual(res["status"], "faerdig", res.get("fejlbesked"))
        self.assertEqual(set_["fil"], r.data)
        self.assertIn("--heat-csv-column", set_["argv"])
        self.assertEqual(res["tidsserier"]["varmelast"]["sha256"], r.sha256)
        shutil.rmtree(tmp, ignore_errors=True)


# ------------------------------------------------------------------ betroede skrivninger

class TestOploesTidsserier(unittest.TestCase):
    def kør(self, svar, valg=None, heat_csv=None, har=True):
        db = StubDb(svar)
        cur = db.forbind().cursor()
        return DB.oploes_tidsserier(cur, UUID2, valg, heat_csv, har), db

    def test_aktive_fryses(self):
        def svar(sql, p):
            if "aktiv=1" in sql:
                return {"rk": [(UUID1.upper(), "a" * 64)]} if p[1] == "varmelast" else {"rk": []}
        ud, db = self.kør(svar)
        self.assertEqual(ud, {"varmelast": {"id": UUID1, "sha256": "a" * 64}})
        self.assertTrue(all("HOLDLOCK" in s for s, _ in db.log))

    def test_to_aktive_er_fejl(self):
        with self.assertRaises(InputFejl):
            self.kør(lambda s, p: {"rk": [(UUID1, "a" * 64), (UUID2, "b" * 64)]} if "aktiv=1" in s else {})

    def test_syntetisk_og_heat_csv_fryser_ikke_aktiv_varmelast(self):
        ud, _ = self.kør(lambda s, p: {"rk": [(UUID1, "a" * 64)]} if p[1] == "varmelast" else {"rk": []},
                         valg={"varmelast": "syntetisk"})
        self.assertEqual(ud, {"varmelast": "syntetisk"})
        ud, db = self.kør(lambda s, p: {"rk": [(UUID1, "a" * 64)]} if p[1] == "varmelast" else {"rk": []},
                          heat_csv="billund_abvaerk_hourly.csv")
        self.assertEqual(ud, {})

    def test_valgt_skal_vaere_vaerkets_gyldige(self):
        for rk in ((UUID1, "a" * 64, "afvist", "varmelast", UUID2),
                   (UUID1, "a" * 64, "gyldig", "solprofil", UUID2),
                   (UUID1, "a" * 64, "gyldig", "varmelast", UUID1)):
            with self.assertRaises(InputFejl, msg=rk):
                self.kør(lambda s, p, rk=rk: {"rk": [rk]} if "WHERE id=?" in s else {"rk": []},
                         valg={"varmelast": UUID1})
        ud, _ = self.kør(lambda s, p: {"rk": [(UUID1, "a" * 64, "gyldig", "varmelast", UUID2)]}
                         if "WHERE id=?" in s else {"rk": []}, valg={"varmelast": UUID1})
        self.assertEqual(ud["varmelast"], {"id": UUID1, "sha256": "a" * 64})

    def test_uden_tabel_intet(self):
        ud, db = self.kør(lambda s, p: {}, har=False)
        self.assertEqual(ud, {})
        self.assertEqual(db.log, [])


class TestRepoTidsserier(unittest.TestCase):
    @unittest.skipUnless(REPO_DATA is not None, "kræver modelrepoets data/")
    def test_casene_registrerer_deres_egne_filer(self):
        cases = REPO_DATA.parent / "cases"
        fund = {}
        for f in sorted(cases.iterdir()):
            if f.suffix == ".yaml" and not f.name.startswith("heat_load"):
                t = f.read_text(encoding="utf-8")
                readme = (REPO_DATA.parent / "README.md").read_text(encoding="utf-8")
                fund[f.stem] = [(x["rolle"], x["fil"], x["aktiv"]) for x in TDB.repo_tidsserier(t, flad(t), f.name, readme)]
        self.assertEqual(fund["andeby"], [("solprofil", "solar_andeby_dmi_karup.csv", True)])
        # Kun den fil, modelrepoet selv kører casen med (casens hoved / README), aldrig en nedarvet kommentar
        # for en anden casefil (runde 8): Spor A kører syntetisk (README.md:176), de rullende med splejset fil.
        splejset = [("varmelast", "billund_abvaerk_hourly_splejset_jun2026.csv", False)]
        self.assertEqual(fund["billund_sporB"], splejset)
        self.assertEqual(fund["billund_idag"], splejset)
        self.assertEqual(fund["billund_sporA_rullende"], splejset)
        self.assertEqual(fund["billund_sporA"], [])
        self.assertEqual(fund["billund_sporA_rullende_bud"], [])
        self.assertEqual(fund["billund_sporB_rullende"], [])

    def test_indsaet_er_idempotent_og_hoejst_en_aktiv(self):
        tmp = Path(tempfile.mkdtemp())
        r = serie(timer())

        def svar_findes(sql, p):
            if "sha256=?" in sql:
                return {"rk": [(UUID1.upper(),)]}
        db = StubDb(svar_findes)
        tid, udfald = TDB.indsaet_tidsserie(db.forbind().cursor(), UUID2, "andeby", "solprofil",
                                            "n", "f.csv", r, "repo", True, "mig", files=str(tmp))
        self.assertEqual((tid, udfald), (UUID1, "uaendret"))
        self.assertFalse(any("INSERT" in s for s, _ in db.log))

        def svar_ny(sql, p):
            if "COUNT(*)" in sql:
                return {"rk": [(1,)]}                      # en anden er allerede aktiv
            if "NEWID()" in sql:
                return {"rk": [(UUID1,)]}
            return {"rk": []}
        db = StubDb(svar_ny)
        TDB.indsaet_tidsserie(db.forbind().cursor(), UUID2, "andeby", "solprofil",
                              "n", "f.csv", r, "repo", True, "mig", files=str(tmp))
        (_, ins), = kald(db, "INSERT INTO dbo.Tidsseries")
        self.assertEqual(ins[18], 0, "ny række må ikke blive aktiv, når en anden er")
        self.assertTrue((tmp / TS.lakehouse_rel("andeby", r.sha256)).is_file())
        shutil.rmtree(tmp)

    def test_validering_json_er_kompakt_og_fejler_hoejlydt(self):
        # Mange huller: oraklet afkorter hullisten selv (huller_fuld=false); resten er kompakt.
        r = serie(timer(doegn=60, spring=set(range(5, 60 * 24, 7))))
        t = TDB.kompakt_validering(r.validering)
        self.assertLessEqual(len(t), 4000)
        self.assertNotIn(": ", t)
        self.assertFalse(json.loads(t)["huller_fuld"])
        with self.assertRaises(ValueError):            # aldrig tavs trunkering (§3.0)
            TDB.kompakt_validering({"advarsler": ["x" * 5000]})


@unittest.skipUnless(HAR_PANDAS and REPO_DATA is not None and (REPO_DATA.parent /
                     "vaerksdata_skabelon.xlsx").is_file(), "kræver modelrepoet, pandas og openpyxl")
class TestVaerksark(unittest.TestCase):
    """Modelrepoets script køres uændret på en udfyldt kopi af skabelonen (8760 timer)."""

    @classmethod
    def setUpClass(cls):
        import openpyxl
        from datetime import datetime, timedelta
        cls.tmp = Path(tempfile.mkdtemp(prefix="fjvvs_"))
        kode = REPO_DATA.parent
        wb = openpyxl.load_workbook(kode / "vaerksdata_skabelon.xlsx")   # i repoets rod fra eef950a (før: doc/)
        ws = wb["Timedata"]
        ws["B3"] = 100
        for rr in range(6, ws.max_row + 1):
            for c in (1, 2, 3):
                ws.cell(rr, c).value = None
        t, rr = datetime(2025, 7, 1), 6
        while t < datetime(2026, 7, 1):
            ws.cell(rr, 1).value = t
            ws.cell(rr, 2).value = 10 + (t.hour % 7) * 0.5
            rr += 1
            t += timedelta(hours=1)
        wb["Priser"]["B30"] = "Test Vaerksark"
        cls.ark = cls.tmp / "ark.xlsx"
        wb.save(cls.ark)
        cls.rc, cls.log = VS.koer_script(kode, sys.executable, cls.ark, cls.tmp / "ud", False)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_konvertering(self):
        self.assertEqual(self.rc, 0, self.log[-1500:])
        k = VS.konverter_output(self.tmp / "ud", "UTC", False)
        self.assertEqual(k.kode, "test_vaerksark")
        prof = [p for p in k.params if p.sti.endswith(".production_profile_path")]
        self.assertEqual([json.loads(p.vaerdi_json) for p in prof], [TS.MARKOER_SOLPROFIL])
        roller = {t["rolle"]: t for t in k.tidsserier}
        self.assertEqual(set(roller), {"varmelast", "solprofil"})
        self.assertEqual(roller["varmelast"]["res"].antal_raekker, 8760)
        self.assertEqual(roller["varmelast"]["res"].start, "2025-07-01T00:00:00Z")
        self.assertIn("syntetisk", roller["solprofil"]["navn"])
        # Runde 8 (V8): oraklets fund om den syntetiske profil siges at gælde modellens profil, ikke arket.
        sol = [a for a in k.advarsler if a.startswith("solprofil")]
        for a in sol:
            self.assertTrue(a.startswith("solprofil (modellens syntetiske profil, ikke jeres data): "), a)
        # Casen kan bygges med begge frosne serier og ingen markør tilbage.
        params = [p.dict() for p in k.params]
        rk = raekker_fra_vaerk(params, (), {
            "model_commit": K.MODEL_COMMIT, "data_commit": K.DF_DATA_SHA,
            "periode_start": "2026-01-05", "periode_slut": "2026-01-06",
            "tidsserier": {r: {"id": UUID1, "sha256": t["res"].sha256} for r, t in roller.items()}})
        plan = byg(rk)
        self.assertNotIn(TS.MARKOER_SOLPROFIL, plan.case_yaml)
        self.assertIn("--heat-csv", plan.cli)

    def test_maalt_solprofil_faar_solvarme_arkets_tidszone(self):
        # Runde 8 (V7): en målt solprofil kommer fra arket Solvarme, som har sin egen B2. Metadata skal være dens,
        # ikke Timedata's. (Talværdierne er scriptets og uændrede; kun Tidsserie.tidszone afhang af fejlen.)
        ud = self.tmp / "maalt"
        shutil.copytree(self.tmp / "ud", ud)
        case = next((ud / "cases").iterdir())
        prof = next(json.loads(p.vaerdi_json) for p in flad(case.read_text(encoding="utf-8"))
                    if p.sti.endswith(".production_profile_path"))
        syn = ud / "data" / Path(prof).name
        maalt = syn.with_name(syn.stem + "_maalt.csv")
        syn.rename(maalt)
        case.write_text(case.read_text(encoding="utf-8").replace(syn.name, maalt.name), encoding="utf-8")
        sol = lambda k: next(t for t in k.tidsserier if t["rolle"] == "solprofil")  # noqa: E731
        k = VS.konverter_output(ud, "UTC", False, "Europe/Copenhagen")
        self.assertIn("målt", sol(k)["navn"])
        self.assertEqual(sol(k)["tidszone"], "Europe/Copenhagen")       # Solvarme!B2, ikke Timedata!B2 (UTC)
        self.assertEqual(sol(VS.konverter_output(ud, "UTC", False))["tidszone"], "UTC")   # uden: som før

    def test_fejl_fra_scriptet_er_laesbar(self):
        log = "Læser x.xlsx\n\nArket kan ikke bruges endnu:\n  Timedata har kun 8 timer. Det ligner\n  eksempelrækkerne.\n\n"
        self.assertEqual(VS.fejl_fra_log(log, 1),
                         "Arket kan ikke bruges endnu: Timedata har kun 8 timer. Det ligner eksempelrækkerne.")

    def test_reserveret_kode_afvises(self):
        ud = self.tmp / "reserv"
        (ud / "cases").mkdir(parents=True)
        (ud / "cases" / "nyt.yaml").write_text("meta:\n  titel: Nyt\n", encoding="utf-8")
        with self.assertRaises(VS.ArkFejl):
            VS.konverter_output(ud, "UTC", False)

    def test_gem_vaerk_i_en_transaktion(self):
        k = VS.konverter_output(self.tmp / "ud", "UTC", False)

        def svar(sql, p):
            if sql.startswith("SELECT status FROM dbo.Vaerksarkimports"):
                return {"rk": [("koerer",)]}
            if "SELECT COUNT(*) FROM dbo.Vaerks" in sql:
                return {"rk": [(0,)]}
            if "NEWID()" in sql:
                return {"rk": [(UUID1,)]}
            if "SELECT COUNT(*) FROM dbo.VaerkParameters" in sql:
                return {"rk": [(len(k.params),)]}
            if "COUNT(*) FROM dbo.Tidsseries" in sql:
                return {"rk": [(0,)]}
            if "UPDATE dbo.Vaerksarkimports" in sql:
                return {"n": 1}
            return {"rk": []}
        db = StubDb(svar)
        files = self.tmp / "Files"
        res = VS.gem_vaerk(db, {"id": UUID2, "filnavn": "ark.xlsx", "uploadet_af": "a@b.dk"}, k,
                           K.MODEL_COMMIT, "log", files=str(files))
        s = [x for x, _ in db.log]
        self.assertEqual(s.count("COMMIT"), 1)
        self.assertLess(max(i for i, x in enumerate(s) if "INSERT INTO dbo.Tidsseries" in x),
                        s.index("COMMIT"))
        self.assertLess(next(i for i, x in enumerate(s) if "SET status='faerdig'" in x), s.index("COMMIT"))
        (ins, _), = kald(db, "INSERT INTO dbo.Vaerks")
        self.assertIn("'app'", ins)
        self.assertEqual(len(res["tidsserier"]), 2)
        for t in res["tidsserier"]:
            self.assertTrue((files / TS.lakehouse_rel("test_vaerksark", t["sha256"])).is_file())

    def test_eksisterende_kode_ruller_tilbage(self):
        k = VS.konverter_output(self.tmp / "ud", "UTC", False)

        def svar(sql, p):
            if sql.startswith("SELECT status FROM dbo.Vaerksarkimports"):
                return {"rk": [("koerer",)]}
            if "SELECT COUNT(*) FROM dbo.Vaerks" in sql:
                return {"rk": [(1,)]}
            return {"rk": []}
        db = StubDb(svar)
        with self.assertRaises(VS.ArkFejl):
            VS.gem_vaerk(db, {"id": UUID2, "uploadet_af": "a"}, k, K.MODEL_COMMIT, "", files=str(self.tmp / "F"))
        self.assertIn("ROLLBACK", [x for x, _ in db.log])
        self.assertNotIn("COMMIT", [x for x, _ in db.log])


if __name__ == "__main__":
    unittest.main()
