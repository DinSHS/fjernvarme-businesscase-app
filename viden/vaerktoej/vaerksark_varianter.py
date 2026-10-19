"""Varianter af doc/vaerksdata_skabelon.xlsx til måling af værksark-kontrakten (fv-data).

Bruges af tjek_data_krav.py --vaerksark. Kræver openpyxl og pandas (fabric/_lokal/.venv).
Hver variant: navn -> (fil, ekstra flag, forventet exitkode, tekst der skal stå i output).
Skriv aldrig varianterne i modelrepoet - de indeholder Billunds målte varmelast.
"""

from __future__ import annotations

import io
from datetime import timedelta
from pathlib import Path


def lav_varianter(skabelon: Path, ud: Path, billund_csv: bytes) -> dict:
    import openpyxl
    import pandas as pd

    ud.mkdir(parents=True, exist_ok=True)
    m = pd.read_csv(io.BytesIO(billund_csv), parse_dates=["timestamp"])
    m = m[(m.timestamp >= "2025-07-01") & (m.timestamp <= "2026-06-30 23:00")]
    par = list(zip(m.timestamp, m.heat_mw_abvaerk))
    gwh = round(m.heat_mw_abvaerk.sum() / 1000, 1)
    lokal = list(zip(m.timestamp.dt.tz_localize("UTC").dt.tz_convert("Europe/Copenhagen").dt.tz_localize(None),
                     m.heat_mw_abvaerk))

    def skriv(navn, *, b3=None, serie=None, tz="UTC"):
        wb = openpyxl.load_workbook(skabelon)
        ws = wb["Timedata"]
        for r in range(6, ws.max_row + 1):
            for c in (1, 2, 3):
                ws.cell(r, c).value = None
        ws["B2"], ws["B3"] = tz, b3
        for i, (t, v) in enumerate(serie or [], start=6):
            ws.cell(i, 1).value = t.to_pydatetime()
            ws.cell(i, 2).value = float(v)
        p = ud / f"{navn}.xlsx"
        wb.save(p)
        return p

    en = timedelta(hours=1)
    return {
        "skabelon_uaendret": (skabelon, [], 1, "Timedata har kun 8 timer"),
        "kun_b3": (skriv("kun_b3", b3=50.0), [], 0, "syntetiseres ud fra årsproduktionen på 50.0 GWh"),
        "maalt_utc": (skriv("maalt_utc", b3=gwh, serie=par), [], 0, "dækning af perioden 1. juli 2025 – 30. juni 2026: 100.0%"),
        "timeslut_uden_flag": (skriv("timeslut", b3=gwh, serie=[(t + en, v) for t, v in par]), [], 1, "Kør igen med --timeslut"),
        "timeslut_med_flag": (ud / "timeslut.xlsx", ["--timeslut"], 0, "læst som timeslut"),
        "dansk_som_utc": (skriv("dansk_som_utc", b3=gwh, serie=lokal), [], 1, "faktisk i dansk tid"),
        "b3_i_mwh": (skriv("b3_mwh", b3=gwh * 1000, serie=par), [], 1, "skal stå i GWh"),
        "hul_30_doegn": (skriv("hul30", b3=gwh, serie=[(t, v) for t, v in par if not (pd.Timestamp("2026-01-01") <= t < pd.Timestamp("2026-01-31"))]),
                         [], 0, "720 timer mangler"),
    }
